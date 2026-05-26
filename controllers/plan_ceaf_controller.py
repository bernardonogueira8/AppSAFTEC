from models.plan_ceaf_model import Plan_ceafModel
from core import os, threading, pd, datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from core.logger import get_logger
from typing import Callable

logger = get_logger("CeafController")

class PlanCeafController:
    """
    Controller for plan_ceaf page
    """

    def __init__(self, model=None):
        self.model = model or Plan_ceafModel
        self.file_path: str = ""
        self.is_processing: bool = False
        self.on_progress: Callable | None = None
        self.on_status:   Callable | None = None
        self.on_done:     Callable | None = None
        self.on_error:    Callable | None = None

    def get_title(self):
        return "Planilhas de Avaliação CEAF"

    def set_file_path(self, path: str):
        self.file_path = path

    def start_processing(self):
        if not self.file_path:
            if self.on_error:
                self.on_error("Nenhum arquivo selecionado.")
            return
        if self.is_processing:
            return
        self.is_processing = True
        threading.Thread(target=self._process, daemon=True).start()

    def _notify_status(self, msg: str):
        logger.info(msg)
        if self.on_status:
            self.on_status(msg)

    def _notify_progress(self, pct: float):
        if self.on_progress:
            self.on_progress(pct)

    def _process(self):
        file_path = self.file_path
        try:
            self._notify_status("Analisando dados...")
            file_name  = os.path.splitext(os.path.basename(file_path))[0]
            output_dir = os.path.join(os.path.dirname(file_path), file_name)
            os.makedirs(output_dir, exist_ok=True)

            xls     = pd.ExcelFile(file_path)
            sheets  = xls.sheet_names
            unidades = set()
            headers  = {}

            for sheet in sheets:
                df = pd.read_excel(xls, sheet_name=sheet, header=None)
                headers[sheet] = df.iloc[0].tolist()
                df.columns = df.iloc[1]
                df = df[2:].reset_index(drop=True)
                if "UNIDADE" in df.columns:
                    unidades.update(df["UNIDADE"].dropna().unique())

            lista_unidades = list(unidades)
            total = len(lista_unidades)

            for i, unidade in enumerate(lista_unidades):
                self._notify_progress((i + 1) / total)
                self._notify_status(f"Processando: {unidade} ({i+1}/{total})")

                output_file = os.path.join(output_dir, f"{unidade}.xlsx")
                with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
                    for sheet in sheets:
                        df = pd.read_excel(xls, sheet_name=sheet, header=1)
                        if "UNIDADE" in df.columns:
                            df_filtered = df[df["UNIDADE"] == unidade]
                            if not df_filtered.empty:
                                header_df    = pd.DataFrame(
                                    [df_filtered.columns.tolist()],
                                    columns=df_filtered.columns,
                                )
                                first_row_df = pd.DataFrame(
                                    [headers[sheet]],
                                    columns=df_filtered.columns,
                                )
                                df_final = pd.concat(
                                    [first_row_df, header_df, df_filtered],
                                    ignore_index=True,
                                )
                                df_final.to_excel(
                                    writer, sheet_name=sheet,
                                    index=False, header=False,
                                )

                wb = load_workbook(output_file)
                for sheet in wb.sheetnames:
                    self._apply_styles(wb[sheet])
                wb.save(output_file)
                logger.info(f"Unidade {unidade} concluída.")

            self._notify_progress(1.0)
            self._notify_status("Concluído!")
            if self.on_done:
                self.on_done(output_dir)

        except Exception as e:
            logger.error(f"Erro no processamento: {e}")
            if self.on_error:
                self.on_error(str(e))
        finally:
            self.is_processing = False

    @staticmethod
    def _apply_styles(ws):
        AZUL_BEBE = "83caff"
        VERMELHO  = "FF0000"
        CINZA     = "dddddd"
        AMARELO   = "FFFF00"
        PRETO     = "000000"

        thin   = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"),  bottom=Side(style="thin"),
        )
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        bold   = Font(bold=True, color=PRETO)
        header = Font(bold=True, color=PRETO, size=12)

        orig_max_row   = ws.max_row
        max_col        = ws.max_column
        antepenult_col = max_col - 1
        footer_row     = orig_max_row + 1

        ws.row_dimensions[1].height = 85
        ws.column_dimensions["A"].width = 5
        if "A1:E1" not in ws.merged_cells:
            ws.merge_cells("A1:E1")
        if max_col >= 6:
            rng = f"F1:{get_column_letter(max_col)}1"
            if rng not in ws.merged_cells:
                ws.merge_cells(rng)

        ws.cell(row=footer_row, column=2).value = "QUANTIDADE A SER LIBERADA POR MÊS"
        ws.cell(row=footer_row, column=6).value = f"=SUM(F3:F{orig_max_row})"
        ws.cell(row=footer_row, column=max_col).value = (
            f"=SUM({get_column_letter(max_col)}3"
            f":{get_column_letter(max_col)}{orig_max_row})"
        )

        for r in range(1, footer_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border    = thin
                cell.alignment = center
                if isinstance(cell.value, datetime):
                    cell.number_format = "DD/MM/YYYY"

                if r == 1:
                    cell.font = header
                    cell.fill = PatternFill(
                        start_color=AZUL_BEBE if c <= 5 else VERMELHO,
                        fill_type="solid",
                    )
                elif r == 2:
                    cell.font  = bold
                    color = AZUL_BEBE if c == max_col else (CINZA if c <= 5 else VERMELHO)
                    cell.fill  = PatternFill(start_color=color, fill_type="solid")
                elif r == footer_row:
                    cell.font = bold
                    if c in [2, 6, max_col]:
                        cell.fill = PatternFill(start_color=CINZA, fill_type="solid")
                elif 2 < r < footer_row:
                    val = str(cell.value).strip() if cell.value is not None else ""
                    if c in [1, 6] and val == "":
                        cell.fill = PatternFill(start_color=AMARELO, fill_type="solid")
                    elif c == antepenult_col and val != "":
                        cell.fill = PatternFill(start_color=AMARELO, fill_type="solid")

        for c in range(2, max_col + 1):
            max_len = max(
                (
                    len(str(ws.cell(row=r, column=c).value))
                    for r in range(1, footer_row + 1)
                    if ws.cell(row=r, column=c).value
                    and not str(ws.cell(row=r, column=c).value).startswith("=")
                ),
                default=0,
            )
            ws.column_dimensions[get_column_letter(c)].width = min(max_len + 5, 50)